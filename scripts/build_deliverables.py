from __future__ import annotations

from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables"
OUT.mkdir(parents=True, exist_ok=True)


def set_run_font(run, name="宋体", size=12, bold=False, color="000000"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9D9D9", size="6"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=140, bottom=100, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_paragraph_border(paragraph, bottom=True):
    p = paragraph._p
    p_pr = p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    if bottom:
        edge = OxmlElement("w:bottom")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "6")
        edge.set(qn("w:space"), "6")
        edge.set(qn("w:color"), "B7B7B7")
        borders.append(edge)


def add_page_field(paragraph):
    run = paragraph.add_run("— ")
    set_run_font(run, "Times New Roman", 10)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, text, fld_end])
    run2 = paragraph.add_run(" —")
    set_run_font(run2, "Times New Roman", 10)


def setup_doc(title: str) -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.35)
    section.right_margin = Cm(2.35)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.45
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.first_line_indent = Cm(0.74)

    for style_name in ("Title", "Heading 1", "Heading 2"):
        style = doc.styles[style_name]
        style.font.name = "黑体" if style_name != "Title" else "黑体"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), style.font.name)
        style.font.color.rgb = RGBColor(0, 0, 0)
    doc.styles["Title"].font.size = Pt(18)
    doc.styles["Title"].font.bold = True
    doc.styles["Heading 1"].font.size = Pt(14)
    doc.styles["Heading 1"].font.bold = True
    doc.styles["Heading 2"].font.size = Pt(12)
    doc.styles["Heading 2"].font.bold = True

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(2)
    hp.paragraph_format.tab_stops.add_tab_stop(Cm(15.8), WD_ALIGN_PARAGRAPH.RIGHT)
    r = hp.add_run("A")
    set_run_font(r, "Arial", 17, True, "225B87")
    r = hp.add_run("I")
    set_run_font(r, "Arial", 17, True, "2AA8C7")
    r = hp.add_run("C")
    set_run_font(r, "Arial", 17, True, "225B87")
    r = hp.add_run("\t第八届全球校园人工智能算法精英大赛 · 算法主题赛")
    set_run_font(r, "宋体", 9, False, "333333")
    set_paragraph_border(hp)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(2)
    add_page_field(fp)

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(14)
    run = p.add_run(title)
    set_run_font(run, "黑体", 18, True)
    return doc


def add_section(doc: Document, text: str):
    p = doc.add_paragraph(style="Heading 1")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    set_run_font(run, "黑体", 14, True)
    return p


def add_subsection(doc: Document, text: str):
    p = doc.add_paragraph(style="Heading 2")
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    set_run_font(run, "黑体", 12, True)
    return p


def add_body(doc: Document, text: str, indent=True, after=5):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = 1.45
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.first_line_indent = Cm(0.74) if indent else Cm(0)
    run = p.add_run(text)
    set_run_font(run, "宋体", 12)
    return p


def add_numbered(doc: Document, n: int, text: str, indent=True):
    return add_body(doc, f"{n}. {text}", indent=indent, after=4)


def add_table(doc: Document, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        c = table.rows[0].cells[i]
        set_cell_shading(c, "2F5D7E")
        set_cell_border(c)
        set_cell_margins(c)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        r = p.add_run(h)
        set_run_font(r, "宋体", 10.5, True, "FFFFFF")
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for i, val in enumerate(row):
            c = cells[i]
            set_cell_border(c)
            set_cell_margins(c)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2:
                set_cell_shading(c, "F3F6F8")
            p = c.paragraphs[0]
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i > 0 else WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(val))
            set_run_font(r, "宋体", 10.5)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def build_video_script():
    doc = setup_doc("工维联演示视频录制流程")
    add_body(doc, "本流程用于录制一段 3 至 5 分钟的工维联真实调用演示视频。视频分为两段：前段录制 https://gongwei.yanyi.online/ 的业务页面，后段录制 https://wt.ioa.pub/ 的梧桐叮当调用。页面返回的方案编号、人员、库存和工单编号以实际结果为准。网络等待时保留状态变化和最终结果画面，可剪去空白等待，但不要剪掉关键状态。")

    add_section(doc, "一、录制前准备")
    add_numbered(doc, 1, "打开 https://gongwei.yanyi.online/，确认首页能加载并停留在主协作页；随后在另一个浏览器标签页打开 https://wt.ioa.pub/。关闭其他标签页和桌面窗口。")
    add_numbered(doc, 2, "使用 OBS Studio 的窗口采集，只选择当前 Edge 窗口；设置 1920×1080 或 1440×900、30 FPS。不要录制密码、验证码、Token、私钥、证书正文或浏览器开发者工具。")
    add_numbered(doc, 3, "演示前准备鼠标和输入法。网页和梧桐叮当都使用下面的三条业务消息；第三条中的方案编号必须复制实际页面显示的编号。")

    add_section(doc, "二、工维联网页流程 00:00 至 03:20")
    add_subsection(doc, "1. 开场 00:00 至 00:20")
    add_body(doc, "画面：停留在主协作页，展示导航和输入区。旁白：工维联用于处理设备异常后的检查协调，系统会先补齐信息，再查询资源，确认后才创建工单。", indent=False)

    add_subsection(doc, "2. 首轮请求 00:20 至 01:00")
    add_body(doc, "在主协作输入框发送：水泵001运行异常，帮我安排明天下午检查。等待页面显示 accepted、working 和补充问题。旁白：系统发现开始时间、预计时长和技能信息仍不完整，因此进入 awaiting_input。此时不要点击创建工单。", indent=False)

    add_subsection(doc, "3. 补充信息与方案 01:00 至 01:40")
    add_body(doc, "在同一对话中发送：下午两点开始，预计一小时，需要泵体检查技能。等待资源查询完成，展示人员、技能、工具、配件、检查时间和待确认方案。此时只展示方案，不进行正式预留，也不要创建工单。", indent=False)

    add_subsection(doc, "4. 确认并建单 01:40 至 02:30")
    add_body(doc, "找到页面显示的真实方案编号，例如 GW-xxxxxxxx；复制该编号后发送：确认方案 GW-xxxxxxxx 并创建工单。把 GW-xxxxxxxx 替换成页面实际编号，不要照抄示例编号。等待资源预留、工单创建和 completed。展示工单编号、设备、时间、人员、工具或配件及当前状态。", indent=False)

    add_subsection(doc, "5. 工单与审计 02:30 至 03:20")
    add_body(doc, "打开工单中心，点开刚刚生成的工单；展示状态、检查时间、人员和物品。再打开资源协同与审计页面，展示资源变化、协作轨迹和操作记录。旁白：这些内容来自现有后端和 SQLite，确认前没有正式工单，确认后只创建一张工单。", indent=False)

    add_section(doc, "三、梧桐叮当真实调用 03:20 至 05:00")
    add_subsection(doc, "1. 选择智能体 03:20 至 03:35")
    add_body(doc, "切换到 https://wt.ioa.pub/。登录会话如果失效，先停止录制，由本人输入密码或验证码，登录完成后再开始本段。进入叮当页面，搜索或点击“工维联核心业务 Leader”，确认右侧当前选中的智能体名称明确显示为“工维联核心业务 Leader”。不要在空白叮当或其他通用智能体中发送任务。", indent=False)

    add_subsection(doc, "2. 梧桐首轮请求 03:35 至 03:55")
    add_body(doc, "在叮当输入框发送：水泵001运行异常，帮我安排明天下午检查。等待 Leader 返回 accepted、working 和 awaiting_input；停留在补充问题页面，录下缺少开始时间、预计时长或技能的提问。此时不要创建工单。", indent=False)

    add_subsection(doc, "3. 梧桐补充信息 03:55 至 04:20")
    add_body(doc, "在同一叮当对话中发送：下午两点开始，预计一小时，需要泵体检查技能。等待页面显示资源协调结果和待确认方案，录下人员、时间、工具、配件和方案编号。此时仍不要确认建单。", indent=False)

    add_subsection(doc, "4. 梧桐确认建单 04:20 至 04:45")
    add_body(doc, "复制叮当页面显示的真实方案编号，例如 GW-xxxxxxxx。将下面消息中的占位符替换成真实编号后发送：确认方案 GW-xxxxxxxx 并创建工单。等待资源 Partner 预留、工单 Partner 创建工单和 completed，录下真实工单编号、设备、时间、人员、工具或配件和当前状态。", indent=False)

    add_subsection(doc, "5. 收尾 04:45 至 05:00")
    add_body(doc, "回到工维联网页或叮当结果页停留 5 秒后停止录制。旁白：网页和梧桐叮当均完成同一条业务链路；声音检测是可选辅助，只提示未发现明显异常或疑似异常；当前平台 HTTPS Access caller 仍待梧桐平台核查，因此 platform_verified 继续为 false。", indent=False)

    add_section(doc, "四、录制后检查")
    add_numbered(doc, 1, "检查网页和梧桐叮当两段画面是否连续，三条消息是否清晰可读，尤其确认确认消息中的方案编号是页面实际编号。")
    add_numbered(doc, 2, "检查视频中没有密码、验证码、Token、私钥、证书正文或其他项目窗口。")
    add_numbered(doc, 3, "导出一个文件：gongwei-main-demo.mp4。视频总时长控制在 3 至 5 分钟。")
    doc.save(OUT / "工维联演示视频录制流程.docx")


def build_technical_report():
    doc = setup_doc("AIC · 工维联智能体互联技术报告")
    add_body(doc, "工维联面向设备异常后的检查协调场景，按照 AIC 与 AIP 的智能体互联思路，将自然语言入口、任务编排、资源协调和工单处理拆分为可追踪的 Leader 与 Partner。本文说明系统背景、需求、技术方案、实施过程、真实验收结果和当前边界，为项目评审、演示视频和后续平台核查提供统一依据。")

    add_section(doc, "一、作品概述")
    add_numbered(doc, 1, "行业与技术背景：设备异常处理通常涉及异常描述、检查条件、人员技能、工具和库存等多类信息。传统单体页面能够记录工单，但难以把自然语言补充、资源协调、人工确认、正式建单和跨平台消息关联在同一条可追踪链路中。工维联采用智能体分工和 AIP 兼容的任务状态，把检查请求转成可验证的协作流程。")
    add_numbered(doc, 2, "作品目标定位：在不迁移既有业务数据和不复制业务动作的前提下，建立 Leader、资源协调 Partner 和工单处理 Partner 的协作系统；支持本机 local_rpc、正式平台 MQ Inbox 和 Leader HTTPS mTLS 三种传输边界；使用 SQLite、幂等、补偿和恢复保证任务、资源和工单一致。")
    add_numbered(doc, 3, "应用边界：人员、库存和设备属于项目已有的模拟演示数据；真实验收对象包括平台网络、AIC 身份、证书、MQ 消息、协议状态、Partner 执行、数据库写入和唯一工单结果。声音检测保留为可插拔辅助工具，只输出“未发现明显异常”或“疑似异常”，不推断具体损坏部件。")

    add_section(doc, "二、需求分析")
    add_numbered(doc, 1, "行业痛点剖析：设备异常请求往往缺少开始时间、预计时长或技能要求；资源信息分散在人员和库存系统；确认前后的动作边界不清会造成重复占用和重复建单；平台调用、MQ 子任务和本地数据库若没有统一编号，难以追踪一次真实协作。")
    add_numbered(doc, 2, "功能需求：系统应接收 text/plain 与 application/json，返回 accepted、working、awaiting_input、awaiting_completion、completed 和 failed 等状态；支持同一业务会话的补充信息；方案必须经过用户确认后才能预留资源和创建工单；重复确认、修改、取消、服务重启和故障补偿均应可验证。")
    add_numbered(doc, 3, "互联需求：Leader 负责任务拆解和确认门，资源 Partner 查询人员、技能、工具和库存，工单 Partner 创建唯一工单；每个 Partner 使用独立 AIC、证书和 Inbox；外部 sessionId、taskId、messageId 与 parent_task_id 进入日志、数据库和审计记录。")
    add_numbered(doc, 4, "安全需求：秘密只放在 secrets 或本机环境变量中，不进入 Git、网页、报告和评审包；正式 Leader 采用 mTLS 和身份绑定；MQ 与 HTTPS 调用不能静默回退到 local_rpc；平台统计未完成时保持 platform_verified=false。")

    add_section(doc, "三、智能体互联技术选型与方案设计")
    add_numbered(doc, 1, "基础开发资源选型：项目使用 Python、SQLite 和官方 ACPs/AIP 兼容接口，保留现有 local_rpc 回归入口。平台通信使用正式 HTTPS Leader 与官方 MQ Inbox Partner，业务动作由共享模块执行，避免为不同传输复制两套业务逻辑。")
    add_numbered(doc, 2, "整体架构设计：网页/API 进入 Leader 编排层，传输选择层在 local_rpc、platform_mq_inbox 和 platform_https_rpc 之间切换，随后调用资源 Partner、工单 Partner 和可选声音 Partner；共享业务动作写入 SQLite 领域与审计层。")
    add_table(doc, ["角色", "职责", "正式入口或传输"], [
        ("Leader", "接收请求、补充信息、生成方案、执行确认门、关联 Partner 结果", "HTTPS JSON-RPC + mTLS"),
        ("资源 Partner", "查询人员技能、工具和配件，执行资源查询与预留", "官方 MQ Inbox"),
        ("工单 Partner", "根据确认后的资源结果创建唯一工单并返回状态", "官方 MQ Inbox"),
        ("声音辅助", "提供异常提示，交给人工确认", "可插拔本机工具"),
    ], widths=[3.2, 8.1, 4.3])
    add_numbered(doc, 3, "关键技术选择：使用官方 Group/Inbox 和 MQ 客户端承载 Partner 消息；使用 SQLite 事务、幂等键和工单编号保护重复操作；使用补偿逻辑在 failed、cancelled 和服务恢复时释放或结算资源；使用受限制的 AMP 元数据转发器记录公开的 AIC、编号、时间、方法和响应码。")
    add_numbered(doc, 4, "方案确认边界：资源查询可以在未确认前执行，正式预留和工单创建必须在用户确认后执行。平台首页会重建外部协议 taskId，项目按明确方案编号、同一 session/caller 校验和持久历史建立业务接续，不把外部 taskId 不变写成已修复。")

    add_section(doc, "四、项目实施")
    add_numbered(doc, 1, "总体实施过程：完成本机 local_rpc 基线、阶段二业务 Partner、阶段五幂等与恢复、正式 MQ Inbox 和 Leader HTTPS mTLS 部署；在现有项目目录 gongwei-step1 内增量修改，不迁移 WIT/Adapter，不安装 Docker，不改其他项目配置。")
    add_numbered(doc, 2, "核心代码实现：Leader 编排位于 agentic_core.py、leader_external_flow.py 和 stage3_leader_aip.py；资源与工单位于 resource_partner_v1.py、work_order_partner_v1.py 和 mq_partner_runtime.py；数据库与领域动作位于 stage2_db.py、stage2_domain.py 和 stage2_business_actions.py；运行入口包括 run_leader_inbox.py、run_partner_inbox.py 和网页服务。")
    add_numbered(doc, 3, "正式身份与部署：三个智能体使用各自 AIC、clientAuth 证书、私钥和 trust bundle；Leader 正式端点为 https://leader.yanyi.online/rpc；资源和工单 Partner 连接 amqps://wt.ioa.pub:5671 的独立 Inbox；秘密文件和服务器备份不进入 Git 或评审压缩包。")
    add_numbered(doc, 4, "测试与证据：本机 369 项测试中 368 项通过、1 项 Windows 符号链接权限跳过；服务器 369/369；阶段二、阶段五、60/60 场景和 24 组网页控件通过。三轮真实叮当调用完成补充信息、待确认方案、正式 MQ 资源预留、唯一工单和 completed。")
    add_numbered(doc, 5, "故障与恢复实现：已验证重复确认不重复创建工单和资源占用，独立修改与取消不覆盖已完成工单，完成工单后工具归还、人员释放和耗材消耗按资源类型结算；共享服务器重启后项目服务、MQ Inbox 和数据库自动恢复。")

    add_section(doc, "五、应用成效")
    add_numbered(doc, 1, "功能成果展示：最新真实会话 session_1841 完成“异常请求→补充信息→资源方案→明确确认→资源预留→工单创建→completed”，工单编号为 WO-20261009-A1044F；重复确认返回同一工单，独立 session_1842 完成改期和取消测试。")
    add_numbered(doc, 2, "协同与追踪效果：主任务与 Partner 子任务能够通过 taskId、messageId、sessionId 和 parent_task_id 对应；确认前累计工单 4 张、active 预留 0 条，确认后累计工单 5 张、active 预留 3 条，重复确认后数量保持不变；SQLite integrity_check=ok。")
    add_numbered(doc, 3, "平台发现与调用：三个 AIC 的默认在线精确查询均返回 HTTP200，未指定名称的叮当自动发现能够推荐工维联核心业务 Leader；平台 Message 统计已显示三套身份记录。平台 HTTPS Access 当前仍为 0，本次捕获到我方 callee 但未找到配对 caller，需梧桐平台核查采集、转发和匹配规则。")
    add_numbered(doc, 4, "性能与运行指标：最新三轮平台成员详情显示首轮约 2.48 秒、第二轮约 11.13 秒、第三轮约 5.41 秒，重复确认约 2.13 秒；Leader 业务状态历史另有独立计时口径，报告不把两种口径混成单一端到端延迟。")
    add_numbered(doc, 5, "关键验证结论：业务调用链已经真实通过，但平台统计的 HTTPS caller 配对仍未完成，因此 platform_verified 继续为 false。不能把网页部署成功、Message 统计或本机回归通过解释为平台全部验收完成。")

    add_section(doc, "六、总结与展望")
    add_numbered(doc, 1, "经验总结：把确认门、幂等、补偿、恢复和跨身份编号作为互联系统的一等约束，可以把一次自然语言任务稳定地落到资源、工单和审计记录；把传输层与业务动作分离，可以在本机回归和正式平台调用之间复用同一套领域逻辑。")
    add_numbered(doc, 2, "当前不足：梧桐平台仍需核查叮当 HTTPS caller 的入库和匹配，项目不能自行生成对端证据；平台客户端会重建外部 taskId，当前方案依靠明确方案编号和调用者校验接续；声音模块只适合辅助演示，不支持具体部件诊断。")
    add_numbered(doc, 3, "后续规划：由平台管理员补齐 HTTPS Access 配对证据后，再复核 discovery、Message 和 Access 的统一时间窗口；随后更新 platform_verified 的最终判定。项目层面继续保持秘密隔离、日志脱敏、数据库完整性和全量回归。")

    add_section(doc, "七、附录")
    add_numbered(doc, 1, "测试数据：系统性能测试原始数据、梧桐平台访问记录、多智能体联调记录、网页控件测试结果、数据库工单记录和资源变化记录，详见《工维联测试数据 20261010.zip》。")
    add_numbered(doc, 2, "完整工程源码：工维联源代码、测试代码、前端源码、运行配置模板和部署脚本，详见源码压缩包。")
    add_numbered(doc, 3, "参考文献：AIP 智能体互联国家标准文件、官方 ACPs SDK 文档、AIP 开源社区技术手册、开发教程及相关学术与行业文献。")
    add_numbered(doc, 4, "证明材料：工维联系统演示视频录制流程、梧桐平台接入操作截图、三轮真实调用截图、系统运行截图和平台统计核查记录。")
    add_numbered(doc, 5, "开源贡献材料：工维联源码压缩包、公开 ACS 文件、AIP 开源社区参考链接、项目测试记录和公开技术说明。")
    doc.save(OUT / "工维联AIC智能体互联技术报告.docx")


def should_include(rel: Path) -> bool:
    parts = {p.lower() for p in rel.parts}
    name = rel.name.lower()
    if any(p in parts for p in {".git", ".venv", ".python-runtime", "secrets", "logs", "output"}):
        return False
    if any(part.startswith("archive") for part in parts):
        return False
    if name in {"gongwei.db", "gongwei.sqlite", "gongwei.sqlite3"}:
        return False
    if name.endswith((".pem", ".key", ".crt", ".p12", ".pfx")):
        return False
    if name.startswith(".env") and name != ".env.example":
        return False
    if rel.parts and rel.parts[0] == "reports":
        return False
    allowed = {".py", ".pyw", ".cmd", ".ps1", ".sh", ".toml", ".lock", ".md", ".json", ".js", ".ts", ".tsx", ".css", ".html", ".txt"}
    return rel.suffix.lower() in allowed


def build_source_zip():
    zip_path = OUT / "gongwei-source-code-20261010.zip"
    manifest = [
        "工维联源码评审包（2026-10-10）",
        "包含：Python 源码、测试、启动脚本、前端源码、公开文档、配置模板和公开 ACS JSON。",
        "排除：secrets、私钥、证书、Token、.env 实例、SQLite/数据库、日志、运行输出、历史评审包和平台截图。",
        "当前 platform_verified=false；本包不代表平台 HTTPS Access caller 已完成匹配。",
        "",
    ]
    files = []
    with ZipFile(zip_path, "w", ZIP_DEFLATED) as zf:
        for path in sorted(ROOT.rglob("*")):
            if not path.is_file():
                continue
            rel = path.relative_to(ROOT)
            if not should_include(rel):
                continue
            arcname = Path("gongwei-step1") / rel
            zf.write(path, arcname.as_posix())
            files.append(arcname.as_posix())
        manifest.extend(files)
        zf.writestr("gongwei-step1/SOURCE_PACKAGE_MANIFEST.txt", "\n".join(manifest) + "\n")
    return zip_path, files


def build_test_data_zip():
    """Package curated, non-secret test evidence for review and submission."""
    zip_path = OUT / "工维联测试数据20261010.zip"
    selected = [
        "reports/stage4-60-scenarios.json",
        "reports/stage4-60-scenarios.md",
        "reports/final-maintenance-audit/unittest-final-result.json",
        "reports/final-maintenance-audit/unittest-final.log",
        "reports/final-maintenance-audit/stage2-final.log",
        "reports/final-maintenance-audit/stage5-final.log",
        "reports/final-maintenance-audit/run_60_scenarios-isolated.log",
        "reports/final-maintenance-audit/browser-awaiting-input.png",
        "reports/final-maintenance-audit/browser-plan.png",
        "reports/final-maintenance-audit/browser-completed.png",
        "reports/final-maintenance-audit/browser-orders.png",
        "reports/final-maintenance-audit/browser-revised-plan.png",
        "reports/frontend-precision-acceptance-20261008/test-results.json",
        "reports/frontend-precision-acceptance-20261008/responsive-metrics.json",
        "reports/frontend-precision-acceptance-20261008/browser-database-final.json",
        "reports/frontend-precision-acceptance-20261008/duplicate-confirm.json",
        "reports/frontend-precision-acceptance-20261008/final-awaiting-input.png",
        "reports/frontend-precision-acceptance-20261008/completed-plan.png",
        "reports/frontend-precision-acceptance-20261008/final-home-1440.png",
        "reports/frontend-precision-acceptance-20261008/final-orders-1440.png",
        "reports/frontend-precision-acceptance-20261008/final-plan-1440.png",
        "reports/frontend-precision-acceptance-20261008/final-resources-1440.png",
        "reports/platform-closeout-20261009/business-flow-assertions.json",
        "reports/platform-closeout-20261009/response-times.json",
        "reports/platform-closeout-20261009/task-message-session-map.csv",
        "reports/platform-closeout-20261009/server-regression-fresh.json",
        "reports/platform-closeout-20261009/server-regression-fresh.log",
        "reports/platform-closeout-20261009/local-regression-latest.log",
        "reports/platform-closeout-20261009/discovery-leader.json",
        "reports/platform-closeout-20261009/discovery-resource_partner.json",
        "reports/platform-closeout-20261009/discovery-workorder_partner.json",
        "reports/platform-closeout-20261009/platform-statistics.json",
        "reports/platform-closeout-20261009/platform-rankings-latest.json",
        "reports/platform-closeout-20261009/amp-access-pairing-investigation.json",
        "reports/platform-closeout-20261009/amp-access-pairing-investigation.md",
        "reports/platform-closeout-20261009/dingdang-round1.txt",
        "reports/platform-closeout-20261009/dingdang-round2.txt",
        "reports/platform-closeout-20261009/dingdang-round3.txt",
        "reports/platform-closeout-20261009/dingdang-completed-detail.txt",
        "reports/platform-closeout-20261009/dingdang-repeat-detail.txt",
        "reports/platform-closeout-20261009/dingdang-auto-discovery.txt",
        "reports/platform-closeout-20261009/dingdang-selected-leader.png",
        "reports/platform-closeout-20261009/dingdang-round1.png",
        "reports/platform-closeout-20261009/dingdang-round2.png",
        "reports/platform-closeout-20261009/dingdang-round3.png",
        "reports/platform-closeout-20261009/dingdang-completed-detail.png",
        "reports/platform-closeout-20261009/dingdang-repeat-confirm.png",
        "reports/platform-closeout-20261009/dingdang-auto-discovery.png",
    ]
    manifest = [
        "工维联测试数据包（2026-10-10）",
        "内容：阶段二、阶段五、60 场景、前端控件、真实叮当三轮、MQ/发现/统计和数据库结果。",
        "用途：项目评审、技术报告附录和演示视频复核。",
        "安全边界：不包含 secrets、私钥、证书、Token、数据库文件或账号密码。",
        "当前平台 HTTPS Access caller 仍待梧桐平台核查，platform_verified=false。",
        "",
    ]
    files = []
    with ZipFile(zip_path, "w", ZIP_DEFLATED) as zf:
        for item in selected:
            path = ROOT / item
            if not path.is_file():
                manifest.append(f"缺失：{item}")
                continue
            arcname = Path("工维联测试数据20261010") / item
            zf.write(path, arcname.as_posix())
            files.append(arcname.as_posix())
        manifest.extend(files)
        zf.writestr("工维联测试数据20261010/TEST_DATA_MANIFEST.txt", "\n".join(manifest) + "\n")
    return zip_path, files


if __name__ == "__main__":
    build_video_script()
    build_technical_report()
    zip_path, files = build_source_zip()
    test_zip, test_files = build_test_data_zip()
    print(f"created={OUT / '工维联演示视频录制流程.docx'}")
    print(f"created={OUT / '工维联AIC智能体互联技术报告.docx'}")
    print(f"created={zip_path} files={len(files)}")
    print(f"created={test_zip} files={len(test_files)}")

